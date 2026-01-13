#!/usr/bin/env python3
"""
SCPI Consolidator V6 - Extraction Avancée avec Techniques Académiques
======================================================================
Version 6.0 avec améliorations basées sur les meilleures pratiques académiques:
- Extraction de tables structurées avec PyMuPDF
- Parsing numérique robuste (K€, M€, Md€)
- Patterns contextuels multi-lignes
- Normalisation et nettoyage des données
- Post-traitement et validation

Conformité : AMF art. 422-227, ASPIM méthodologie 01/01/2022, ANC 2016-03
"""

import re
import os
import sys
from dataclasses import dataclass, field, asdict
from typing import Optional, List, Dict, Tuple, Any
from datetime import datetime
from pathlib import Path
from decimal import Decimal, InvalidOperation
import warnings

import pandas as pd
import fitz  # PyMuPDF for PDF parsing
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, NamedStyle
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.formatting.rule import FormulaRule, ColorScaleRule, CellIsRule
from openpyxl.utils import get_column_letter
from openpyxl.chart import BarChart, PieChart, Reference
from openpyxl.chart.label import DataLabelList

warnings.filterwarnings('ignore')


# =============================================================================
# DATACLASSES - Structures de données (identiques à V5)
# =============================================================================

@dataclass
class IndicateursSCPI:
    """90+ indicateurs réglementaires SCPI"""
    source_pdf: str = ""
    date_extraction: str = ""
    qualite_extraction: str = ""
    nb_indicateurs_extraits: int = 0

    # Identification
    nom_scpi: Optional[str] = None
    societe_gestion: Optional[str] = None
    type_scpi: Optional[str] = None
    visa_amf: Optional[str] = None
    date_creation: Optional[str] = None

    # Capital et Parts
    capital_nominal: Optional[float] = None
    capital_statutaire_max: Optional[float] = None
    nombre_parts: Optional[int] = None
    nombre_associes: Optional[int] = None
    capitalisation: Optional[float] = None
    valeur_nominale: Optional[float] = None

    # Prix et Valeurs
    prix_souscription: Optional[float] = None
    prix_retrait: Optional[float] = None
    prime_emission: Optional[float] = None
    valeur_realisation: Optional[float] = None
    valeur_realisation_par_part: Optional[float] = None
    valeur_reconstitution: Optional[float] = None
    valeur_reconstitution_par_part: Optional[float] = None
    valeur_venale_patrimoine: Optional[float] = None
    valeur_ifi_resident: Optional[float] = None
    valeur_ifi_non_resident: Optional[float] = None
    ecart_prix_reconstitution: Optional[float] = None

    # Performance ASPIM
    taux_distribution: Optional[float] = None
    dividende_brut: Optional[float] = None
    dividende_exceptionnel: Optional[float] = None
    report_a_nouveau: Optional[float] = None
    report_a_nouveau_par_part: Optional[float] = None
    tri_5_ans: Optional[float] = None
    tri_10_ans: Optional[float] = None
    tri_15_ans: Optional[float] = None
    tri_20_ans: Optional[float] = None
    rendement_global_immobilier: Optional[float] = None

    # Occupation
    tof_annuel: Optional[float] = None
    tof_t1: Optional[float] = None
    tof_t2: Optional[float] = None
    tof_t3: Optional[float] = None
    tof_t4: Optional[float] = None
    top_annuel: Optional[float] = None
    surface_vacante: Optional[float] = None
    surface_vacante_pct: Optional[float] = None

    # Patrimoine
    nombre_actifs: Optional[int] = None
    nombre_actifs_direct: Optional[int] = None
    nombre_actifs_indirect: Optional[int] = None
    surface_totale: Optional[float] = None
    prix_moyen_m2: Optional[float] = None

    # Répartition typologique
    pct_bureaux: Optional[float] = None
    pct_commerces: Optional[float] = None
    pct_logistique: Optional[float] = None
    pct_activites: Optional[float] = None
    pct_sante: Optional[float] = None
    pct_residentiel: Optional[float] = None
    pct_hotellerie: Optional[float] = None
    pct_enseignement: Optional[float] = None
    pct_autres: Optional[float] = None

    # Répartition géographique
    pct_paris: Optional[float] = None
    pct_idf_hors_paris: Optional[float] = None
    pct_regions: Optional[float] = None
    pct_etranger: Optional[float] = None
    pct_france: Optional[float] = None

    # Acquisitions / Cessions
    nb_acquisitions: Optional[int] = None
    montant_acquisitions: Optional[float] = None
    nb_cessions: Optional[int] = None
    montant_cessions: Optional[float] = None
    plus_value_cessions: Optional[float] = None
    moins_value_cessions: Optional[float] = None

    # Collecte
    collecte_brute: Optional[float] = None
    collecte_nette: Optional[float] = None
    parts_en_attente_retrait: Optional[int] = None
    montant_attente_retrait: Optional[float] = None
    delai_jouissance: Optional[str] = None

    # Endettement
    ratio_endettement_ltv: Optional[float] = None
    emprunts_bancaires: Optional[float] = None
    dettes_financieres: Optional[float] = None
    tresorerie: Optional[float] = None
    endettement_net: Optional[float] = None

    # Résultats comptables
    loyers_quittances: Optional[float] = None
    produits_locatifs: Optional[float] = None
    produits_financiers: Optional[float] = None
    charges_immobilieres: Optional[float] = None
    resultat_immobilier: Optional[float] = None
    resultat_net: Optional[float] = None
    resultat_par_part: Optional[float] = None
    benefice_comptable: Optional[float] = None
    revenus_fonciers_par_part: Optional[float] = None
    revenus_financiers_par_part: Optional[float] = None

    # Frais
    commission_souscription_pct: Optional[float] = None
    commission_gestion_pct: Optional[float] = None
    commission_arbitrage_pct: Optional[float] = None
    commission_souscription_montant: Optional[float] = None
    commission_gestion_montant: Optional[float] = None
    provision_gros_entretien: Optional[float] = None

    # CAPEX / Travaux
    total_travaux: Optional[float] = None
    travaux_amelioration: Optional[float] = None
    travaux_gros_entretien: Optional[float] = None
    travaux_renovation: Optional[float] = None
    travaux_mise_conformite: Optional[float] = None
    travaux_environnementaux: Optional[float] = None
    capex_par_m2: Optional[float] = None
    ratio_capex_valeur_venale: Optional[float] = None


@dataclass
class ActifImmobilier:
    """Actif immobilier détenu en direct"""
    scpi_source: str = ""
    id_actif: str = ""
    nom_actif: Optional[str] = None
    adresse: Optional[str] = None
    ville: Optional[str] = None
    code_postal: Optional[str] = None
    pays: str = "France"
    region: Optional[str] = None
    type_actif: Optional[str] = None
    date_acquisition: Optional[str] = None
    surface_m2: Optional[float] = None
    nb_lots: Optional[int] = None
    quote_part_pct: float = 100.0
    valeur_venale: Optional[float] = None
    prix_acquisition: Optional[float] = None
    loyer_annuel: Optional[float] = None
    taux_occupation: Optional[float] = None
    nb_locataires: Optional[int] = None
    date_fin_bail_principale: Optional[str] = None
    travaux_total: float = 0.0
    travaux_renovation: float = 0.0
    travaux_gros_entretien: float = 0.0
    travaux_mise_conformite: float = 0.0
    travaux_environnementaux: float = 0.0
    travaux_amelioration: float = 0.0
    travaux_autres: float = 0.0
    annee_travaux: Optional[int] = None
    description_travaux: Optional[str] = None
    statut_travaux: Optional[str] = None


@dataclass
class SCISousJacente:
    """SCI ou OPCI détenue par la SCPI"""
    scpi_source: str = ""
    nom_sci: Optional[str] = None
    type_structure: Optional[str] = None
    affectation: Optional[str] = None
    date_acquisition: Optional[str] = None
    date_constitution: Optional[str] = None
    quote_part_detention: Optional[float] = None
    nombre_immeubles: Optional[int] = None
    surface_m2: Optional[float] = None
    valeur_venale_hd: Optional[float] = None
    droits: Optional[float] = None
    prix_acquisition: Optional[float] = None
    travaux: Optional[float] = None
    valeur_nette_comptable: Optional[float] = None
    valeur_estimee: Optional[float] = None
    ecart_acquisition: Optional[float] = None
    immeubles: Optional[float] = None
    dettes: Optional[float] = None


@dataclass
class CAPEXDetail:
    """Détail CAPEX par actif"""
    scpi_source: str = ""
    id_actif: str = ""
    nom_actif: Optional[str] = None
    type_actif: Optional[str] = None
    ville: Optional[str] = None
    surface_m2: Optional[float] = None
    valeur_venale: Optional[float] = None
    travaux_total: float = 0.0
    travaux_renovation: float = 0.0
    travaux_gros_entretien: float = 0.0
    travaux_mise_conformite: float = 0.0
    travaux_environnementaux: float = 0.0
    travaux_amelioration: float = 0.0
    travaux_autres: float = 0.0
    capex_par_m2: Optional[float] = None
    capex_sur_vv_pct: Optional[float] = None
    annee_travaux: Optional[int] = None
    description_travaux: Optional[str] = None
    statut_travaux: Optional[str] = None
    source_donnee: str = "Actif Direct"


@dataclass
class CAPEXSynthese:
    """Synthèse CAPEX par SCPI"""
    scpi_source: str = ""
    total_capex: float = 0.0
    capex_actifs_directs: float = 0.0
    capex_sci: float = 0.0
    nb_actifs_total: int = 0
    nb_actifs_avec_capex: int = 0
    pct_actifs_avec_capex: float = 0.0
    nb_sci_total: int = 0
    nb_sci_avec_capex: int = 0
    capex_moyen_par_actif: float = 0.0
    capex_median_par_actif: float = 0.0
    capex_min: float = 0.0
    capex_max: float = 0.0
    capex_sur_vv_pct: Optional[float] = None
    capex_par_m2: Optional[float] = None
    capex_sur_loyers_pct: Optional[float] = None
    pct_renovation: float = 0.0
    pct_gros_entretien: float = 0.0
    pct_mise_conformite: float = 0.0
    pct_environnemental: float = 0.0
    pct_amelioration: float = 0.0
    pct_autres: float = 0.0
    surface_totale_concernee: float = 0.0


@dataclass
class TestCoherence:
    """Résultat d'un test de cohérence"""
    scpi_source: str = ""
    categorie: str = ""
    test_id: str = ""
    description: str = ""
    valeur_testee: Optional[str] = None
    valeur_reference: Optional[str] = None
    ecart_pct: Optional[float] = None
    statut: str = ""
    commentaire: str = ""
    severite: str = "Info"


# =============================================================================
# CLASSE PRINCIPALE - Consolidateur V6 avec extraction avancée
# =============================================================================

class SCPIConsolidatorV6:
    """Consolidateur SCPI V6 avec techniques d'extraction académiques avancées"""

    # Sociétés de gestion connues (référentiel étendu)
    KNOWN_MANAGERS = [
        'Sofidy', 'Amundi', 'Amundi Immobilier', 'BNP Paribas REIM', 'BNP Paribas Real Estate',
        'La Française', 'La Française REM', 'Primonial REIM', 'Perial AM', 'PERIAL Asset Management',
        'AEW Patrimoine', 'AEW Ciloger', 'AEW', 'Swiss Life AM', 'Swiss Life Asset Managers',
        'Corum AM', 'Corum Asset Management', 'Paref Gestion', 'Inter Gestion REIM', 'Inter Gestion',
        'Advenis REIM', 'Advenis', 'HSBC REIM', 'HSBC Real Estate', 'Alderan',
        'Altarea', 'Altarea Investment Managers', 'Altixia REIM', 'Altixia',
        'Allianz Real Estate', 'Allianz', 'Novaxia Investissement', 'Novaxia',
        'Atland Voisin', 'Atland', 'Iroko', 'Iroko Zen', 'Remake AM', 'Remake',
        'Euryale AM', 'Euryale', 'Norma Capital', 'Epsilon Capital', 'Praemia REIM',
        'Sogenial Immobilier', 'Fiducial Gérance', 'Voisin', 'Kyaneos AM',
        'EDMOND DE ROTHSCHILD REIM', 'Edmond de Rothschild'
    ]

    # Multiplicateurs pour parsing numérique
    MULTIPLIERS = {
        'md€': 1_000_000_000, 'md €': 1_000_000_000, 'milliard': 1_000_000_000,
        'm€': 1_000_000, 'm €': 1_000_000, 'million': 1_000_000, 'millions': 1_000_000,
        'k€': 1_000, 'k €': 1_000, 'millier': 1_000, 'milliers': 1_000,
        '€': 1, 'euro': 1, 'euros': 1
    }

    def __init__(self, output_dir: str = "output", input_dir: str = "input"):
        self.output_dir = Path(output_dir)
        self.output_dir.mkdir(parents=True, exist_ok=True)
        self.input_dir = Path(input_dir)
        if not self.input_dir.exists():
            self.input_dir.mkdir(parents=True, exist_ok=True)

        self.pdf_files = list(self.input_dir.glob("*.pdf")) + list(self.input_dir.glob("*.PDF"))

        self.results = {
            'indicateurs': [],
            'actifs': [],
            'sci': [],
            'capex_detail': [],
            'capex_synthese': [],
            'controles': [],
            'erreurs': []
        }

        self.seuils = {
            'ecart_valeur_venale_pct': 5.0,
            'capex_max_pct_valeur': 20.0,
            'capex_normal_min_m2': 30.0,
            'capex_normal_max_m2': 250.0,
            'tof_min': 50.0,
            'tof_max': 100.0,
            'td_min': 0.0,
            'td_max': 15.0,
            'ltv_max': 50.0,
            'somme_geo_tolerance': 2.0,
            'somme_typo_tolerance': 2.0,
        }

        self._init_styles()

    def _init_styles(self):
        """Initialise les styles Excel"""
        self.styles = {
            'header': {
                'fill': PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid"),
                'font': Font(color="FFFFFF", bold=True, size=11),
                'alignment': Alignment(horizontal='center', vertical='center', wrap_text=True)
            },
            'ok': PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid"),
            'warning': PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid"),
            'error': PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid"),
            'total': {
                'fill': PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid"),
                'font': Font(bold=True, size=11)
            },
            'border': Border(
                left=Side(style='thin', color='B4B4B4'),
                right=Side(style='thin', color='B4B4B4'),
                top=Side(style='thin', color='B4B4B4'),
                bottom=Side(style='thin', color='B4B4B4')
            )
        }

    # =========================================================================
    # PARSING NUMÉRIQUE AVANCÉ (Technique académique #1)
    # =========================================================================

    def _normalize_number(self, text: str) -> str:
        """Normalise un texte numérique en supprimant espaces et formatage"""
        if not text:
            return ""
        # Supprimer espaces insécables, espaces normaux, etc.
        text = text.replace('\u202f', '').replace('\xa0', '').replace(' ', '')
        # Remplacer virgule décimale par point
        text = text.replace(',', '.')
        # Supprimer caractères non numériques sauf point et moins
        text = re.sub(r'[^\d.\-]', '', text)
        return text

    def _parse_amount(self, text: str) -> Optional[float]:
        """
        Parse un montant avec gestion intelligente des unités (K€, M€, Md€).
        Technique académique: parsing contextuel avec multiplicateurs.
        """
        if not text:
            return None

        # Patterns pour extraire nombre + unité
        patterns = [
            # Format: "123,45 M€" ou "123.45 M €"
            r'([\d\s,\.]+)\s*(Md€|Md\s*€|M€|M\s*€|K€|K\s*€|€|millions?|milliards?|milliers?)',
            # Format: nombre seul
            r'([\d\s,\.]+)',
        ]

        for pattern in patterns:
            match = re.search(pattern, text, re.IGNORECASE)
            if match:
                num_str = self._normalize_number(match.group(1))
                if not num_str:
                    continue

                try:
                    value = float(num_str)
                except ValueError:
                    continue

                # Déterminer le multiplicateur
                unit = match.group(2).lower() if len(match.groups()) > 1 and match.group(2) else '€'
                for key, mult in self.MULTIPLIERS.items():
                    if key in unit:
                        value *= mult
                        break

                return value

        return None

    def _extract_number_advanced(self, text: str, patterns: List[str],
                                  min_val: Optional[float] = None,
                                  max_val: Optional[float] = None) -> Optional[float]:
        """
        Extraction numérique avancée avec validation de plage.
        Technique académique: multi-pattern matching avec validation.
        """
        for pattern in patterns:
            matches = re.finditer(pattern, text, re.IGNORECASE | re.MULTILINE)
            for match in matches:
                try:
                    # Chercher le groupe capturant
                    raw_value = match.group(1) if match.lastindex else match.group(0)
                    value = self._parse_amount(raw_value)

                    if value is not None:
                        # Validation de plage
                        if min_val is not None and value < min_val:
                            continue
                        if max_val is not None and value > max_val:
                            continue
                        return value
                except (ValueError, AttributeError, IndexError):
                    continue

        return None

    def _extract_percentage_advanced(self, text: str, patterns: List[str],
                                      min_val: float = 0.0,
                                      max_val: float = 100.0) -> Optional[float]:
        """
        Extraction de pourcentage avancée avec validation.
        """
        for pattern in patterns:
            matches = re.finditer(pattern, text, re.IGNORECASE | re.MULTILINE)
            for match in matches:
                try:
                    raw_value = match.group(1)
                    value = float(self._normalize_number(raw_value))

                    if min_val <= value <= max_val:
                        return value
                except (ValueError, AttributeError, IndexError):
                    continue

        return None

    # =========================================================================
    # EXTRACTION DE TABLES (Technique académique #2)
    # =========================================================================

    def _extract_tables_from_pdf(self, doc: fitz.Document) -> List[Dict]:
        """
        Extrait les données structurées des tables du PDF.
        Technique académique: extraction de tables avec PyMuPDF.
        """
        all_tables_data = []

        for page_num, page in enumerate(doc):
            try:
                tables = page.find_tables()
                for table in tables:
                    if table.row_count > 1:  # Au moins en-tête + 1 ligne
                        table_data = {
                            'page': page_num + 1,
                            'rows': [],
                            'headers': []
                        }

                        for row_idx, row in enumerate(table.extract()):
                            cleaned_row = [str(cell).strip() if cell else "" for cell in row]
                            if row_idx == 0:
                                table_data['headers'] = cleaned_row
                            else:
                                table_data['rows'].append(cleaned_row)

                        all_tables_data.append(table_data)
            except Exception:
                continue

        return all_tables_data

    def _search_in_tables(self, tables: List[Dict], key_patterns: List[str]) -> Optional[str]:
        """
        Recherche une valeur dans les tables extraites.
        """
        for table in tables:
            for row in table['rows']:
                for cell_idx, cell in enumerate(row):
                    for pattern in key_patterns:
                        if re.search(pattern, cell, re.IGNORECASE):
                            # Retourner la cellule suivante si disponible
                            if cell_idx + 1 < len(row) and row[cell_idx + 1]:
                                return row[cell_idx + 1]
        return None

    # =========================================================================
    # EXTRACTION CONTEXTUELLE (Technique académique #3)
    # =========================================================================

    def _extract_with_context(self, text: str, anchor_pattern: str,
                               value_pattern: str, context_size: int = 200) -> Optional[str]:
        """
        Extraction contextuelle: trouve une ancre puis cherche la valeur dans le contexte.
        Technique académique: extraction basée sur le contexte.
        """
        matches = list(re.finditer(anchor_pattern, text, re.IGNORECASE))

        for match in matches:
            start = max(0, match.start() - context_size)
            end = min(len(text), match.end() + context_size)
            context = text[start:end]

            value_match = re.search(value_pattern, context, re.IGNORECASE)
            if value_match:
                return value_match.group(1) if value_match.lastindex else value_match.group(0)

        return None

    # =========================================================================
    # PARSING PDF PRINCIPAL
    # =========================================================================

    def _parse_pdf(self, pdf_path: Path) -> Optional[IndicateursSCPI]:
        """
        Parse un PDF avec techniques d'extraction avancées.
        """
        try:
            doc = fitz.open(pdf_path)
            full_text = ""

            # Extraction du texte complet
            for page in doc:
                full_text += page.get_text() + "\n"

            # Extraction des tables structurées
            tables = self._extract_tables_from_pdf(doc)
            doc.close()

            # Créer l'objet indicateurs
            indicateurs = IndicateursSCPI(
                source_pdf=pdf_path.name,
                date_extraction=datetime.now().strftime("%Y-%m-%d %H:%M:%S")
            )

            # =====================================================================
            # IDENTIFICATION
            # =====================================================================

            # Nom SCPI (depuis fichier + validation texte)
            indicateurs.nom_scpi = self._extract_scpi_name(pdf_path, full_text)

            # Société de gestion (référentiel + patterns)
            indicateurs.societe_gestion = self._extract_management_company(full_text)

            # Visa AMF
            visa_patterns = [
                r'[Vv]isa\s+(?:AMF|de\s+l[\'\']AMF)[^\d]*n[°o]?\s*(\d{2}[-/]\d+)',
                r'[Vv]isa\s+(?:AMF)?[^\d]*(\d{2}[-/]\d{2,5})',
                r'AMF[^\d]{0,20}(\d{2}[-/]\d+)',
            ]
            for pattern in visa_patterns:
                match = re.search(pattern, full_text[:10000])
                if match:
                    indicateurs.visa_amf = match.group(1)
                    break

            # Date de création
            date_patterns = [
                r'[Cc]r[ée][ée]e?\s+(?:en\s+|le\s+)?(\d{1,2}[/\-]\d{1,2}[/\-]\d{4})',
                r'[Cc]r[ée][ée]e?\s+(?:en\s+)?(\d{4})',
                r'[Dd]ate\s+de\s+cr[ée]ation[^\d]{0,20}(\d{4})',
                r'[Cc]onstitu[ée]e?\s+(?:en\s+)?(\d{4})',
            ]
            for pattern in date_patterns:
                match = re.search(pattern, full_text[:20000])
                if match:
                    indicateurs.date_creation = match.group(1)
                    break

            # Type SCPI
            indicateurs.type_scpi = self._detect_scpi_type(full_text)

            # =====================================================================
            # CAPITAL ET PARTS
            # =====================================================================

            # Capitalisation
            cap_patterns = [
                r'[Cc]apitalisation[^\d]{0,30}([\d\s,\.]+)\s*(?:M€|M\s*€|millions?)',
                r'[Cc]apitalisation[^\d]{0,30}([\d\s,\.]+)\s*(?:€|euros?)',
                r'capitalisation\s+(?:au\s+31/12/2024)?[^\d]{0,30}([\d\s,\.]+)',
            ]
            cap = self._extract_number_advanced(full_text, cap_patterns, min_val=1000)
            if cap:
                # Normaliser en euros
                if cap < 100000:  # Probablement en M€
                    cap *= 1_000_000
                indicateurs.capitalisation = cap

            # Recherche aussi dans les tables
            if not indicateurs.capitalisation:
                table_val = self._search_in_tables(tables, [r'capitalisation'])
                if table_val:
                    indicateurs.capitalisation = self._parse_amount(table_val)

            # Nombre d'associés
            assoc_patterns = [
                r'[Nn]ombre\s+d[\'\']associ[ée]s[^\d]{0,20}([\d\s]+)',
                r'([\d\s]+)\s+associ[ée]s',
                r'associ[ée]s[^\d]{0,10}([\d\s]+)',
            ]
            assoc = self._extract_number_advanced(full_text, assoc_patterns, min_val=10, max_val=500000)
            if assoc:
                indicateurs.nombre_associes = int(assoc)

            # Nombre de parts
            parts_patterns = [
                r'[Nn]ombre\s+(?:total\s+)?de\s+parts?[^\d]{0,20}([\d\s]+)',
                r'([\d\s]+)\s+parts?\s+(?:en\s+circulation|émises?)',
                r'parts?\s+(?:sociales?)?[^\d]{0,10}([\d\s]+)',
            ]
            parts = self._extract_number_advanced(full_text, parts_patterns, min_val=1000, max_val=100000000)
            if parts:
                indicateurs.nombre_parts = int(parts)

            # Capital nominal
            cn_patterns = [
                r'[Cc]apital\s+(?:social\s+)?nominal[^\d]{0,20}([\d\s,\.]+)\s*€',
                r'[Cc]apital\s+social[^\d]{0,20}([\d\s,\.]+)\s*€',
            ]
            indicateurs.capital_nominal = self._extract_number_advanced(full_text, cn_patterns, min_val=1000)

            # Valeur nominale
            vn_patterns = [
                r'[Vv]aleur\s+nominale[^\d]{0,20}([\d,\.]+)\s*€',
                r'[Nn]ominal(?:e)?\s+(?:de\s+)?(?:la\s+)?part[^\d]{0,20}([\d,\.]+)\s*€',
            ]
            indicateurs.valeur_nominale = self._extract_number_advanced(full_text, vn_patterns, min_val=1, max_val=5000)

            # =====================================================================
            # PRIX ET VALEURS
            # =====================================================================

            # Prix de souscription
            ps_patterns = [
                r'[Pp]rix\s+(?:de\s+)?souscription[^\d]{0,30}([\d\s,\.]+)\s*€',
                r'souscription[^\d]{0,20}([\d,\.]+)\s*€\s*/\s*part',
                r'souscription[^\d]{0,30}([\d,\.]+)\s*€',
            ]
            indicateurs.prix_souscription = self._extract_number_advanced(full_text, ps_patterns, min_val=50, max_val=10000)

            # Prix de retrait
            pr_patterns = [
                r'[Pp]rix\s+(?:de\s+)?retrait[^\d]{0,30}([\d\s,\.]+)\s*€',
                r'retrait[^\d]{0,20}([\d,\.]+)\s*€\s*/\s*part',
            ]
            indicateurs.prix_retrait = self._extract_number_advanced(full_text, pr_patterns, min_val=50, max_val=10000)

            # Valeur de réalisation
            vr_patterns = [
                r'[Vv]aleur\s+de\s+r[ée]alisation[^\d]{0,30}([\d\s,\.]+)\s*€',
                r'r[ée]alisation[^\d]{0,20}([\d,\.]+)\s*€\s*/\s*part',
            ]
            indicateurs.valeur_realisation = self._extract_number_advanced(full_text, vr_patterns, min_val=50, max_val=10000)

            # Valeur de reconstitution
            vrec_patterns = [
                r'[Vv]aleur\s+de\s+reconstitution[^\d]{0,30}([\d\s,\.]+)\s*€',
                r'reconstitution[^\d]{0,20}([\d,\.]+)\s*€\s*/\s*part',
            ]
            indicateurs.valeur_reconstitution = self._extract_number_advanced(full_text, vrec_patterns, min_val=50, max_val=10000)

            # Valeur vénale du patrimoine
            vvp_patterns = [
                r'[Vv]aleur\s+v[ée]nale\s+(?:du\s+)?(?:patrimoine|total)[^\d]{0,30}([\d\s,\.]+)\s*(?:M€|M\s*€)',
                r'[Vv]aleur\s+v[ée]nale[^\d]{0,30}([\d\s,\.]+)\s*(?:M€|millions?)',
                r'patrimoine[^\d]{0,30}([\d\s,\.]+)\s*(?:M€|M\s*€)',
            ]
            vvp = self._extract_number_advanced(full_text, vvp_patterns, min_val=1)
            if vvp:
                if vvp < 100000:
                    vvp *= 1_000_000
                indicateurs.valeur_venale_patrimoine = vvp

            # IFI
            ifi_patterns = [
                r'[Vv]aleur\s+IFI[^\d]{0,20}([\d,\.]+)\s*€',
                r'IFI[^\d]{0,20}([\d,\.]+)\s*€\s*/?\s*part',
            ]
            indicateurs.valeur_ifi_resident = self._extract_number_advanced(full_text, ifi_patterns, min_val=1, max_val=5000)

            # =====================================================================
            # PERFORMANCE ASPIM
            # =====================================================================

            # Taux de distribution (TD)
            td_patterns = [
                r'[Tt]aux\s+de\s+distribution[^\d]{0,30}([\d,\.]+)\s*%',
                r'TD\s*(?:2024|2023)?[^\d]{0,10}([\d,\.]+)\s*%',
                r'distribution[^\d]{0,20}([\d,\.]+)\s*%\s*(?:brut|net)?',
            ]
            indicateurs.taux_distribution = self._extract_percentage_advanced(full_text, td_patterns, min_val=0, max_val=15)

            # Dividende brut
            div_patterns = [
                r'[Dd]ividende\s+(?:brut\s+)?(?:par\s+part)?[^\d]{0,20}([\d,\.]+)\s*€',
                r'distribution\s+(?:par\s+part)?[^\d]{0,20}([\d,\.]+)\s*€',
            ]
            indicateurs.dividende_brut = self._extract_number_advanced(full_text, div_patterns, min_val=1, max_val=500)

            # TRI (5, 10, 15, 20 ans)
            tri_configs = [
                ('tri_5_ans', [r'TRI\s*(?:à\s*)?5\s*ans?[^\d]{0,10}([\-\d,\.]+)\s*%']),
                ('tri_10_ans', [r'TRI\s*(?:à\s*)?10\s*ans?[^\d]{0,10}([\-\d,\.]+)\s*%']),
                ('tri_15_ans', [r'TRI\s*(?:à\s*)?15\s*ans?[^\d]{0,10}([\-\d,\.]+)\s*%']),
                ('tri_20_ans', [r'TRI\s*(?:à\s*)?20\s*ans?[^\d]{0,10}([\-\d,\.]+)\s*%']),
            ]
            for attr, patterns in tri_configs:
                value = self._extract_percentage_advanced(full_text, patterns, min_val=-20, max_val=20)
                if value:
                    setattr(indicateurs, attr, value)

            # Rendement global immobilier
            rgi_patterns = [
                r'[Rr]endement\s+global\s+immobilier[^\d]{0,20}([\-\d,\.]+)\s*%',
                r'RGI[^\d]{0,10}([\-\d,\.]+)\s*%',
            ]
            indicateurs.rendement_global_immobilier = self._extract_percentage_advanced(
                full_text, rgi_patterns, min_val=-20, max_val=20
            )

            # =====================================================================
            # OCCUPATION
            # =====================================================================

            # TOF annuel
            tof_patterns = [
                r'[Tt]aux\s+d[\'\']occupation\s+financier[^\d]{0,20}([\d,\.]+)\s*%',
                r'TOF\s*(?:moyen\s+)?(?:annuel)?[^\d]{0,10}([\d,\.]+)\s*%',
                r'occupation\s+financi[eè]re[^\d]{0,20}([\d,\.]+)\s*%',
            ]
            indicateurs.tof_annuel = self._extract_percentage_advanced(full_text, tof_patterns, min_val=50, max_val=100)

            # TOP annuel
            top_patterns = [
                r'[Tt]aux\s+d[\'\']occupation\s+physique[^\d]{0,20}([\d,\.]+)\s*%',
                r'TOP[^\d]{0,10}([\d,\.]+)\s*%',
            ]
            indicateurs.top_annuel = self._extract_percentage_advanced(full_text, top_patterns, min_val=50, max_val=100)

            # TOF trimestriels
            for q, attr in [(1, 'tof_t1'), (2, 'tof_t2'), (3, 'tof_t3'), (4, 'tof_t4')]:
                patterns = [
                    rf'T{q}[^\d]{{0,20}}([\d,\.]+)\s*%[^\n]{{0,30}}(?:TOF|occupation)',
                    rf'(?:TOF|occupation)[^\n]{{0,30}}T{q}[^\d]{{0,10}}([\d,\.]+)\s*%',
                ]
                value = self._extract_percentage_advanced(full_text, patterns, min_val=50, max_val=100)
                if value:
                    setattr(indicateurs, attr, value)

            # Surface vacante
            sv_patterns = [
                r'[Ss]urface\s+vacante[^\d]{0,20}([\d\s,\.]+)\s*m[²2]',
                r'vacance[^\d]{0,20}([\d\s,\.]+)\s*m[²2]',
            ]
            indicateurs.surface_vacante = self._extract_number_advanced(full_text, sv_patterns, min_val=0)

            # =====================================================================
            # PATRIMOINE
            # =====================================================================

            # Nombre d'actifs
            na_patterns = [
                r'[Nn]ombre\s+d[\'\'](?:actifs?|immeubles?)[^\d]{0,20}(\d+)',
                r'(\d+)\s+(?:actifs?|immeubles?)\s+(?:au|en|à)',
                r'portefeuille[^\d]{0,30}(\d+)\s+(?:actifs?|immeubles?)',
                r'compos[ée]\s+de\s+(\d+)\s+(?:actifs?|immeubles?)',
            ]
            na = self._extract_number_advanced(full_text, na_patterns, min_val=1, max_val=500)
            if na:
                indicateurs.nombre_actifs = int(na)

            # Surface totale
            st_patterns = [
                r'[Ss]urface\s+(?:totale|globale)[^\d]{0,20}([\d\s,\.]+)\s*m[²2]',
                r'([\d\s]+)\s*m[²2]\s+(?:de\s+surface|au\s+total)',
            ]
            indicateurs.surface_totale = self._extract_number_advanced(full_text, st_patterns, min_val=100)

            # =====================================================================
            # RÉPARTITION GÉOGRAPHIQUE
            # =====================================================================

            geo_configs = [
                ('pct_paris', [r'Paris[^\d]{0,30}([\d,\.]+)\s*%', r'([\d,\.]+)\s*%[^\n]{0,20}Paris']),
                ('pct_idf_hors_paris', [r'[ÎI]le[- ]de[- ]France[^\d]{0,30}([\d,\.]+)\s*%', r'IDF[^\d]{0,10}([\d,\.]+)\s*%']),
                ('pct_regions', [r'[Rr][ée]gions?[^\d]{0,30}([\d,\.]+)\s*%', r'[Pp]rovince[^\d]{0,30}([\d,\.]+)\s*%']),
                ('pct_etranger', [r'[ÉE]tranger[^\d]{0,30}([\d,\.]+)\s*%', r'[Ii]nternational[^\d]{0,30}([\d,\.]+)\s*%']),
                ('pct_france', [r'France[^\d]{0,30}([\d,\.]+)\s*%']),
            ]
            for attr, patterns in geo_configs:
                value = self._extract_percentage_advanced(full_text, patterns, min_val=0, max_val=100)
                if value:
                    setattr(indicateurs, attr, value)

            # =====================================================================
            # RÉPARTITION TYPOLOGIQUE
            # =====================================================================

            typo_configs = [
                ('pct_bureaux', [r'[Bb]ureaux[^\d]{0,30}([\d,\.]+)\s*%']),
                ('pct_commerces', [r'[Cc]ommerces?[^\d]{0,30}([\d,\.]+)\s*%']),
                ('pct_logistique', [r'[Ll]ogistique[^\d]{0,30}([\d,\.]+)\s*%', r'[Ee]ntrep[ôo]ts?[^\d]{0,30}([\d,\.]+)\s*%']),
                ('pct_activites', [r'[Aa]ctivit[ée]s?[^\d]{0,30}([\d,\.]+)\s*%', r'[Ll]ocaux\s+d[\'\']activit[ée][^\d]{0,20}([\d,\.]+)\s*%']),
                ('pct_sante', [r'[Ss]ant[ée][^\d]{0,30}([\d,\.]+)\s*%', r'[Mm][ée]dical[^\d]{0,30}([\d,\.]+)\s*%']),
                ('pct_residentiel', [r'[Rr][ée]sidentiel[^\d]{0,30}([\d,\.]+)\s*%', r'[Hh]abitation[^\d]{0,30}([\d,\.]+)\s*%']),
                ('pct_hotellerie', [r'[Hh][ôo]tel(?:s|lerie)?[^\d]{0,30}([\d,\.]+)\s*%']),
                ('pct_enseignement', [r'[Ee]nseignement[^\d]{0,30}([\d,\.]+)\s*%', r'[ÉE]ducation[^\d]{0,30}([\d,\.]+)\s*%']),
            ]
            for attr, patterns in typo_configs:
                value = self._extract_percentage_advanced(full_text, patterns, min_val=0, max_val=100)
                if value:
                    setattr(indicateurs, attr, value)

            # =====================================================================
            # ACQUISITIONS / CESSIONS
            # =====================================================================

            # Montant acquisitions
            acq_patterns = [
                r'[Aa]cquisitions?[^\d]{0,30}([\d\s,\.]+)\s*(?:M€|M\s*€|millions?)',
                r'([\d\s,\.]+)\s*(?:M€|M\s*€)\s+d[\'\']acquisitions?',
            ]
            acq = self._extract_number_advanced(full_text, acq_patterns, min_val=0.1)
            if acq:
                if acq < 10000:
                    acq *= 1_000_000
                indicateurs.montant_acquisitions = acq

            # Nombre d'acquisitions
            nb_acq_patterns = [
                r'(\d+)\s+acquisitions?',
                r'acquisitions?[^\d]{0,10}(\d+)',
            ]
            nb_acq = self._extract_number_advanced(full_text, nb_acq_patterns, min_val=0, max_val=100)
            if nb_acq:
                indicateurs.nb_acquisitions = int(nb_acq)

            # Montant cessions
            cess_patterns = [
                r'[Cc]essions?[^\d]{0,30}([\d\s,\.]+)\s*(?:M€|M\s*€|millions?)',
                r'([\d\s,\.]+)\s*(?:M€|M\s*€)\s+de\s+cessions?',
            ]
            cess = self._extract_number_advanced(full_text, cess_patterns, min_val=0.1)
            if cess:
                if cess < 10000:
                    cess *= 1_000_000
                indicateurs.montant_cessions = cess

            # Plus-values de cession
            pv_patterns = [
                r'[Pp]lus[- ]values?\s+(?:de\s+)?cessions?[^\d]{0,20}([\d\s,\.]+)\s*(?:€|K€|M€)',
                r'[Pp]lus[- ]values?[^\d]{0,30}([\d\s,\.]+)\s*(?:€|M€)',
            ]
            indicateurs.plus_value_cessions = self._extract_number_advanced(full_text, pv_patterns, min_val=0)

            # =====================================================================
            # COLLECTE
            # =====================================================================

            # Collecte brute
            cb_patterns = [
                r'[Cc]ollecte\s+brute[^\d]{0,30}([\d\s,\.]+)\s*(?:M€|M\s*€|millions?)',
                r'[Cc]ollecte\s+brute[^\d]{0,30}([\d\s,\.]+)\s*€',
            ]
            cb = self._extract_number_advanced(full_text, cb_patterns, min_val=0)
            if cb:
                if cb < 100000:
                    cb *= 1_000_000
                indicateurs.collecte_brute = cb

            # Collecte nette
            cn_patterns = [
                r'[Cc]ollecte\s+nette[^\d]{0,30}([\d\s,\.]+)\s*(?:M€|M\s*€|millions?)',
            ]
            cn = self._extract_number_advanced(full_text, cn_patterns, min_val=0)
            if cn:
                if cn < 100000:
                    cn *= 1_000_000
                indicateurs.collecte_nette = cn

            # Délai de jouissance
            dj_patterns = [
                r'd[ée]lai\s+de\s+jouissance[^\d]{0,20}(\d+)\s*mois',
                r'jouissance[^\d]{0,20}(\d+)\s*mois',
            ]
            for pattern in dj_patterns:
                match = re.search(pattern, full_text, re.IGNORECASE)
                if match:
                    indicateurs.delai_jouissance = f"{match.group(1)} mois"
                    break

            # =====================================================================
            # ENDETTEMENT
            # =====================================================================

            # LTV
            ltv_patterns = [
                r'LTV[^\d]{0,20}([\d,\.]+)\s*%',
                r'[Rr]atio\s+d[\'\']endettement[^\d]{0,20}([\d,\.]+)\s*%',
                r'[Ee]ndettement[^\d]{0,30}([\d,\.]+)\s*%',
            ]
            indicateurs.ratio_endettement_ltv = self._extract_percentage_advanced(
                full_text, ltv_patterns, min_val=0, max_val=60
            )

            # Emprunts bancaires
            emp_patterns = [
                r'[Ee]mprunts?\s+bancaires?[^\d]{0,30}([\d\s,\.]+)\s*(?:M€|M\s*€|€)',
                r'[Dd]ette\s+bancaire[^\d]{0,30}([\d\s,\.]+)\s*(?:M€|M\s*€)',
            ]
            emp = self._extract_number_advanced(full_text, emp_patterns, min_val=0)
            if emp:
                if emp < 100000:
                    emp *= 1_000_000
                indicateurs.emprunts_bancaires = emp

            # Trésorerie
            tres_patterns = [
                r'[Tt]r[ée]sorerie[^\d]{0,30}([\d\s,\.]+)\s*(?:M€|M\s*€|€)',
            ]
            tres = self._extract_number_advanced(full_text, tres_patterns, min_val=0)
            if tres:
                if tres < 100000:
                    tres *= 1_000_000
                indicateurs.tresorerie = tres

            # =====================================================================
            # RÉSULTATS COMPTABLES
            # =====================================================================

            # Loyers
            loyers_patterns = [
                r'[Ll]oyers?\s+(?:quittanc[ée]s?|encaiss[ée]s?)[^\d]{0,30}([\d\s,\.]+)\s*(?:M€|K€|€)',
                r'[Pp]roduits\s+locatifs[^\d]{0,30}([\d\s,\.]+)\s*(?:M€|K€|€)',
            ]
            loyers = self._extract_number_advanced(full_text, loyers_patterns, min_val=0)
            if loyers:
                if loyers < 100000:
                    loyers *= 1_000_000
                indicateurs.loyers_quittances = loyers

            # Résultat net
            rn_patterns = [
                r'[Rr][ée]sultat\s+net[^\d]{0,30}([\d\s,\.]+)\s*(?:M€|K€|€)',
            ]
            rn = self._extract_number_advanced(full_text, rn_patterns, min_val=0)
            if rn:
                if rn < 100000:
                    rn *= 1_000_000
                indicateurs.resultat_net = rn

            # Résultat par part
            rpp_patterns = [
                r'[Rr][ée]sultat\s+(?:net\s+)?par\s+part[^\d]{0,20}([\d,\.]+)\s*€',
            ]
            indicateurs.resultat_par_part = self._extract_number_advanced(full_text, rpp_patterns, min_val=0, max_val=500)

            # =====================================================================
            # FRAIS ET COMMISSIONS
            # =====================================================================

            # Commission de souscription
            cs_patterns = [
                r'[Cc]ommission\s+(?:de\s+)?souscription[^\d]{0,30}([\d,\.]+)\s*%',
                r'[Ff]rais\s+(?:de\s+)?souscription[^\d]{0,30}([\d,\.]+)\s*%',
            ]
            indicateurs.commission_souscription_pct = self._extract_percentage_advanced(
                full_text, cs_patterns, min_val=0, max_val=15
            )

            # Commission de gestion
            cg_patterns = [
                r'[Cc]ommission\s+de\s+gestion[^\d]{0,30}([\d,\.]+)\s*%',
                r'[Ff]rais\s+de\s+gestion[^\d]{0,30}([\d,\.]+)\s*%',
            ]
            indicateurs.commission_gestion_pct = self._extract_percentage_advanced(
                full_text, cg_patterns, min_val=0, max_val=20
            )

            # Commission d'arbitrage
            ca_patterns = [
                r'[Cc]ommission\s+d[\'\']arbitrage[^\d]{0,30}([\d,\.]+)\s*%',
            ]
            indicateurs.commission_arbitrage_pct = self._extract_percentage_advanced(
                full_text, ca_patterns, min_val=0, max_val=10
            )

            # =====================================================================
            # TRAVAUX / CAPEX
            # =====================================================================

            # Total travaux
            trav_patterns = [
                r'[Tt]ravaux[^\d]{0,30}([\d\s,\.]+)\s*(?:M€|K€|€)',
                r'CAPEX[^\d]{0,30}([\d\s,\.]+)\s*(?:M€|K€|€)',
            ]
            trav = self._extract_number_advanced(full_text, trav_patterns, min_val=0)
            if trav:
                indicateurs.total_travaux = trav

            # Provision gros entretien
            pge_patterns = [
                r'[Pp]rovision\s+(?:pour\s+)?gros\s+entretien[^\d]{0,30}([\d\s,\.]+)\s*€',
            ]
            indicateurs.provision_gros_entretien = self._extract_number_advanced(full_text, pge_patterns, min_val=0)

            # =====================================================================
            # EXTRACTION DES ACTIFS ET SCIs
            # =====================================================================

            self._extract_actifs(full_text, indicateurs.nom_scpi)
            self._extract_scis(full_text, indicateurs.nom_scpi)

            # =====================================================================
            # CALCULS DÉRIVÉS
            # =====================================================================

            self._compute_derived_values(indicateurs)

            # =====================================================================
            # COMPTAGE ET QUALITÉ
            # =====================================================================

            count = sum(1 for field_name in indicateurs.__dataclass_fields__
                       if getattr(indicateurs, field_name) not in (None, "", 0))
            indicateurs.nb_indicateurs_extraits = count

            indicateurs.qualite_extraction = (
                "Excellente" if count >= 40 else
                "Très bonne" if count >= 30 else
                "Bonne" if count >= 20 else
                "Moyenne" if count >= 10 else "Faible"
            )

            return indicateurs

        except Exception as e:
            print(f"   Erreur parsing {pdf_path.name}: {e}")
            import traceback
            traceback.print_exc()
            return None

    def _extract_scpi_name(self, pdf_path: Path, full_text: str) -> str:
        """Extrait le nom de la SCPI de manière robuste"""
        # D'abord depuis le nom de fichier
        pdf_name = pdf_path.stem
        clean_name = pdf_name

        # Nettoyer le nom du fichier
        clean_name = re.sub(r'_RA_?\d{4}', '', clean_name)
        clean_name = re.sub(r'-RA-?\d{4}', '', clean_name)
        clean_name = re.sub(r'Rapport\s*annuel\s*\d{4}', '', clean_name, flags=re.IGNORECASE)
        clean_name = re.sub(r'France\s*SCPI\s*[-_]?\s*', '', clean_name, flags=re.IGNORECASE)
        clean_name = re.sub(r'^9328[-_]?\s*', '', clean_name)
        clean_name = re.sub(r'[-_]+', ' ', clean_name)
        clean_name = re.sub(r'\s+', ' ', clean_name).strip()

        # Si le nom est trop court, chercher dans le texte
        if len(clean_name) < 5:
            name_patterns = [
                r'SCPI\s+([A-Z][A-Za-zÀ-ÿ\s\-\']{3,40})',
                r'([A-Z][A-Z\s\-]+(?:PATRIMOINE|PIERRE|PLACEMENT|IMMOBILIER))',
            ]
            for pattern in name_patterns:
                match = re.search(pattern, full_text[:5000])
                if match:
                    candidate = match.group(1).strip()
                    if 5 < len(candidate) < 50:
                        clean_name = candidate
                        break

        return clean_name[:60] if clean_name else pdf_path.stem[:50]

    def _extract_management_company(self, full_text: str) -> Optional[str]:
        """Extrait la société de gestion"""
        # D'abord chercher les noms connus
        for manager in self.KNOWN_MANAGERS:
            if re.search(rf'\b{re.escape(manager)}\b', full_text, re.IGNORECASE):
                return manager

        # Puis utiliser des patterns génériques
        sg_patterns = [
            r'[Ss]oci[ée]t[ée]\s+de\s+gestion[^\n]{0,50}([A-Z][A-Za-z\s\-]+(?:REIM|AM|Gestion))',
            r'[Gg][ée]r[ée]e?\s+par\s+([A-Z][A-Za-z\s\-]+)',
        ]
        for pattern in sg_patterns:
            match = re.search(pattern, full_text[:50000])
            if match:
                sg = match.group(1).strip()
                if 3 < len(sg) < 50:
                    return sg

        return None

    def _detect_scpi_type(self, full_text: str) -> Optional[str]:
        """Détecte le type de SCPI"""
        text_lower = full_text[:20000].lower()

        if 'diversifi' in text_lower:
            return "Diversifiée"
        elif 'bureaux' in text_lower and 'commerce' not in text_lower:
            return "Bureaux"
        elif 'commerce' in text_lower and 'bureaux' not in text_lower:
            return "Commerces"
        elif 'logistique' in text_lower or 'activit' in text_lower:
            return "Logistique/Activités"
        elif 'santé' in text_lower or 'medical' in text_lower:
            return "Santé"
        elif 'résidentiel' in text_lower or 'habitation' in text_lower:
            return "Résidentiel"
        elif 'hôtel' in text_lower:
            return "Hôtellerie"

        return "Diversifiée"

    def _compute_derived_values(self, ind: IndicateursSCPI):
        """Calcule les valeurs dérivées"""
        # Prix moyen au m²
        if ind.valeur_venale_patrimoine and ind.surface_totale and ind.surface_totale > 0:
            ind.prix_moyen_m2 = ind.valeur_venale_patrimoine / ind.surface_totale

        # Écart prix/reconstitution
        if ind.prix_souscription and ind.valeur_reconstitution and ind.valeur_reconstitution > 0:
            ind.ecart_prix_reconstitution = (
                (ind.prix_souscription - ind.valeur_reconstitution) / ind.valeur_reconstitution * 100
            )

        # Endettement net
        if ind.emprunts_bancaires and ind.tresorerie:
            ind.endettement_net = ind.emprunts_bancaires - ind.tresorerie

        # Commissions montants
        if ind.commission_souscription_pct and ind.prix_souscription:
            ind.commission_souscription_montant = (ind.commission_souscription_pct / 100) * ind.prix_souscription

        # CAPEX par m²
        if ind.total_travaux and ind.surface_totale and ind.surface_totale > 0:
            ind.capex_par_m2 = ind.total_travaux / ind.surface_totale

        # Ratio CAPEX / Valeur vénale
        if ind.total_travaux and ind.valeur_venale_patrimoine and ind.valeur_venale_patrimoine > 0:
            ind.ratio_capex_valeur_venale = (ind.total_travaux / ind.valeur_venale_patrimoine) * 100

    def _extract_actifs(self, text: str, scpi_source: str):
        """Extrait les actifs immobiliers"""
        # Pattern pour adresses françaises
        actif_pattern = r'(\d+[^€\d\n]{3,80})\s+(\d{5})\s+([A-ZÀÂÉÈÊËÏÎÔÙÛÜ][A-ZÀÂÉÈÊËÏÎÔÙÛÜ\s\-]{2,30})'

        matches = list(re.finditer(actif_pattern, text))
        seen_addresses = set()

        for match in matches:
            adresse_raw = match.group(1).strip()
            code_postal = match.group(2)
            ville = match.group(3).strip()

            # Nettoyer l'adresse
            adresse = re.sub(r'\s*\([^)]*%[^)]*\)\s*$', '', adresse_raw).strip()

            # Vérifier que c'est une vraie adresse
            if not re.search(r'(rue|avenue|boulevard|quai|place|allée|impasse|chemin|passage|cours|voie)', adresse, re.IGNORECASE):
                continue

            if len(adresse) < 10 or len(ville) < 2:
                continue

            key = f"{adresse.lower()}_{code_postal}"
            if key in seen_addresses:
                continue
            seen_addresses.add(key)

            # Contexte pour extraction
            start = max(0, match.start() - 20)
            end = min(len(text), match.end() + 400)
            context = text[start:end]

            actif = ActifImmobilier(
                scpi_source=scpi_source,
                adresse=adresse,
                code_postal=code_postal,
                ville=ville.strip()
            )

            # Quote-part
            quote_match = re.search(r'(\d+)\s*%\s*(?:de\s+l[\'\']indivision)?', adresse_raw)
            if quote_match:
                actif.quote_part_pct = float(quote_match.group(1))

            # Surface
            surface_match = re.search(r'(\d[\d\s]{0,8})\s*(?:m²|m2)', context)
            if surface_match:
                try:
                    val = surface_match.group(1).replace(' ', '').replace('\u202f', '')
                    actif.surface_m2 = float(val)
                except:
                    pass

            # Valeur vénale
            valeur_match = re.search(r'(\d{1,3}(?:[\s\u202f]\d{3})+(?:,\d+)?)\s*€', context)
            if valeur_match:
                try:
                    val = valeur_match.group(1).replace(' ', '').replace('\u202f', '').replace(',', '.')
                    actif.valeur_venale = float(val)
                except:
                    pass

            # Type d'actif
            type_match = re.search(r'\b(Bureaux|Commerces?|Logistique|Activit[ée]s?|H[ôo]tels?|R[ée]sidentiel|Sant[ée])\b', context, re.IGNORECASE)
            if type_match:
                actif.type_actif = type_match.group(1).capitalize()

            self.results['actifs'].append(actif)

    def _extract_scis(self, text: str, scpi_source: str):
        """Extrait les SCIs détenues"""
        sci_pattern = r'(SCI|OPPCI|OPCI)\s+([A-Z][A-Z\s\-0-9]+?)(?:\s+Titres|\s+\d|\s*$)'

        matches = re.finditer(sci_pattern, text)
        seen = set()

        for match in matches:
            type_struct = match.group(1)
            nom = match.group(2).strip()

            if nom in seen or len(nom) < 3:
                continue
            seen.add(nom)

            sci = SCISousJacente(
                scpi_source=scpi_source,
                nom_sci=nom,
                type_structure=type_struct
            )

            # Contexte
            sci_section = text[max(0, match.start()-100):min(len(text), match.end()+500)]

            # Surface
            surface_match = re.search(r'(\d[\d\s]*)\s*m[²2]', sci_section)
            if surface_match:
                try:
                    sci.surface_m2 = float(surface_match.group(1).replace(' ', ''))
                except:
                    pass

            self.results['sci'].append(sci)

    # =========================================================================
    # CONSOLIDATION ET EXPORT
    # =========================================================================

    def consolidate(self):
        """Effectue toutes les consolidations"""
        if self.pdf_files:
            self._load_pdfs()

        print("🔧 Consolidation des CAPEX...")
        self._build_capex_synthese()

        print("✅ Tests de cohérence...")
        self._run_all_tests()

        return self.results

    def _load_pdfs(self):
        """Charge et parse tous les PDFs"""
        print(f"📄 Extraction des données depuis {len(self.pdf_files)} PDF(s)...")

        for pdf_path in self.pdf_files:
            print(f"   📖 {pdf_path.name}")
            indicateurs = self._parse_pdf(pdf_path)
            if indicateurs:
                self.results['indicateurs'].append(indicateurs)
                print(f"      ✓ {indicateurs.nb_indicateurs_extraits} indicateurs ({indicateurs.qualite_extraction})")
            else:
                self.results['erreurs'].append(f"Échec: {pdf_path.name}")

    def _build_capex_synthese(self):
        """Construit la synthèse CAPEX par SCPI"""
        for ind in self.results['indicateurs']:
            scpi = ind.nom_scpi or ind.source_pdf

            synthese = CAPEXSynthese(scpi_source=scpi)

            actifs_scpi = [a for a in self.results['actifs'] if a.scpi_source == scpi]
            sci_scpi = [s for s in self.results['sci'] if s.scpi_source == scpi]

            synthese.nb_actifs_total = len(actifs_scpi)
            synthese.nb_sci_total = len(sci_scpi)

            if ind.total_travaux:
                synthese.total_capex = ind.total_travaux

            if ind.valeur_venale_patrimoine and ind.valeur_venale_patrimoine > 0 and synthese.total_capex:
                synthese.capex_sur_vv_pct = (synthese.total_capex / ind.valeur_venale_patrimoine) * 100

            if ind.surface_totale and ind.surface_totale > 0 and synthese.total_capex:
                synthese.capex_par_m2 = synthese.total_capex / ind.surface_totale

            self.results['capex_synthese'].append(synthese)

    def _run_all_tests(self):
        """Exécute tous les tests de cohérence"""
        for ind in self.results['indicateurs']:
            scpi = ind.nom_scpi or ind.source_pdf

            # Test TD
            if ind.taux_distribution:
                statut = "OK" if 2 <= ind.taux_distribution <= 12 else "ALERTE"
                self.results['controles'].append(TestCoherence(
                    scpi_source=scpi,
                    categorie="PERFORMANCE",
                    test_id="PERF-001",
                    description="TD dans plage normale (2-12%)",
                    valeur_testee=f"{ind.taux_distribution:.2f}%",
                    statut=statut
                ))

            # Test TOF
            if ind.tof_annuel:
                statut = "OK" if 80 <= ind.tof_annuel <= 100 else "ALERTE"
                self.results['controles'].append(TestCoherence(
                    scpi_source=scpi,
                    categorie="OCCUPATION",
                    test_id="OCC-001",
                    description="TOF dans plage normale (80-100%)",
                    valeur_testee=f"{ind.tof_annuel:.2f}%",
                    statut=statut
                ))

            # Test LTV
            if ind.ratio_endettement_ltv:
                statut = "OK" if ind.ratio_endettement_ltv <= 40 else "ALERTE"
                self.results['controles'].append(TestCoherence(
                    scpi_source=scpi,
                    categorie="ENDETTEMENT",
                    test_id="END-001",
                    description="LTV < 40%",
                    valeur_testee=f"{ind.ratio_endettement_ltv:.2f}%",
                    statut=statut
                ))

            # Test cohérence géographique
            geo_values = [v for v in [ind.pct_paris, ind.pct_idf_hors_paris, ind.pct_regions, ind.pct_etranger] if v]
            if len(geo_values) >= 2:
                total = sum(geo_values)
                statut = "OK" if 95 <= total <= 105 else "ALERTE"
                self.results['controles'].append(TestCoherence(
                    scpi_source=scpi,
                    categorie="REPARTITION",
                    test_id="REP-001",
                    description="Somme géographique ≈ 100%",
                    valeur_testee=f"{total:.1f}%",
                    statut=statut
                ))

    def export_excel(self, filename: str = "scpi_consolidation_v6.xlsx") -> str:
        """Export vers Excel"""
        wb = Workbook()

        # Onglet Indicateurs
        ws = wb.active
        ws.title = "Indicateurs"

        if self.results['indicateurs']:
            df = pd.DataFrame([asdict(d) for d in self.results['indicateurs']])

            # En-têtes
            for col, header in enumerate(df.columns, 1):
                cell = ws.cell(row=1, column=col, value=header)
                cell.fill = self.styles['header']['fill']
                cell.font = self.styles['header']['font']
                cell.alignment = self.styles['header']['alignment']

            # Données
            for row_idx, row_data in enumerate(df.values, 2):
                for col_idx, value in enumerate(row_data, 1):
                    ws.cell(row=row_idx, column=col_idx, value=value)

        # Onglet Patrimoine
        if self.results['actifs']:
            ws_actifs = wb.create_sheet("Patrimoine")
            df_actifs = pd.DataFrame([asdict(d) for d in self.results['actifs']])
            for col, header in enumerate(df_actifs.columns, 1):
                cell = ws_actifs.cell(row=1, column=col, value=header)
                cell.fill = self.styles['header']['fill']
                cell.font = self.styles['header']['font']
            for row_idx, row_data in enumerate(df_actifs.values, 2):
                for col_idx, value in enumerate(row_data, 1):
                    ws_actifs.cell(row=row_idx, column=col_idx, value=value)

        # Onglet SCI
        if self.results['sci']:
            ws_sci = wb.create_sheet("SCI")
            df_sci = pd.DataFrame([asdict(d) for d in self.results['sci']])
            for col, header in enumerate(df_sci.columns, 1):
                cell = ws_sci.cell(row=1, column=col, value=header)
                cell.fill = self.styles['header']['fill']
                cell.font = self.styles['header']['font']
            for row_idx, row_data in enumerate(df_sci.values, 2):
                for col_idx, value in enumerate(row_data, 1):
                    ws_sci.cell(row=row_idx, column=col_idx, value=value)

        # Onglet Contrôles
        if self.results['controles']:
            ws_ctrl = wb.create_sheet("Controles")
            df_ctrl = pd.DataFrame([asdict(d) for d in self.results['controles']])
            for col, header in enumerate(df_ctrl.columns, 1):
                cell = ws_ctrl.cell(row=1, column=col, value=header)
                cell.fill = self.styles['header']['fill']
                cell.font = self.styles['header']['font']
            for row_idx, row_data in enumerate(df_ctrl.values, 2):
                for col_idx, value in enumerate(row_data, 1):
                    cell = ws_ctrl.cell(row=row_idx, column=col_idx, value=value)
                    # Coloration statut
                    if df_ctrl.columns[col_idx-1] == 'statut':
                        if value == "OK":
                            cell.fill = self.styles['ok']
                        elif value == "ALERTE":
                            cell.fill = self.styles['warning']

        # Sauvegarder
        path = self.output_dir / filename
        wb.save(path)
        return str(path)

    def export_csv(self, prefix: str = "scpi_2024"):
        """Export CSV"""
        outputs = []

        exports = [
            ('indicateurs', self.results['indicateurs']),
            ('patrimoine', self.results['actifs']),
            ('sci', self.results['sci']),
            ('capex_synthese', self.results['capex_synthese']),
            ('controles', self.results['controles']),
        ]

        for name, data in exports:
            if data:
                df = pd.DataFrame([asdict(d) for d in data])

                # Nettoyage des colonnes vides
                df = df.dropna(axis=1, how='all')

                path = self.output_dir / f"{prefix}_{name}_v6.csv"
                df.to_csv(path, index=False, encoding='utf-8-sig')
                outputs.append(str(path))

        return outputs

    def print_summary(self):
        """Affiche le résumé"""
        print(f"\n{'='*80}")
        print("CONSOLIDATION SCPI V6 - EXTRACTION ACADÉMIQUE AVANCÉE")
        print(f"{'='*80}")

        print(f"\n📊 VOLUMES DE DONNÉES")
        print(f"   SCPI analysées:     {len(self.results['indicateurs']):>5}")
        print(f"   Actifs directs:     {len(self.results['actifs']):>5}")
        print(f"   SCI sous-jacentes:  {len(self.results['sci']):>5}")
        print(f"   Tests de cohérence: {len(self.results['controles']):>5}")

        # Statistiques d'extraction
        if self.results['indicateurs']:
            avg_indicators = sum(i.nb_indicateurs_extraits for i in self.results['indicateurs']) / len(self.results['indicateurs'])
            max_indicators = max(i.nb_indicateurs_extraits for i in self.results['indicateurs'])
            print(f"\n📈 QUALITÉ D'EXTRACTION")
            print(f"   Indicateurs moyens par SCPI: {avg_indicators:.1f}")
            print(f"   Maximum indicateurs:         {max_indicators}")

            # Taux de remplissage
            total_fields = len(IndicateursSCPI.__dataclass_fields__) - 4  # Sans metadata
            fill_rate = (avg_indicators / total_fields) * 100
            print(f"   Taux de remplissage moyen:   {fill_rate:.1f}%")

        # Contrôles
        if self.results['controles']:
            nb_ok = len([c for c in self.results['controles'] if c.statut == "OK"])
            nb_alerte = len([c for c in self.results['controles'] if c.statut == "ALERTE"])
            print(f"\n✅ CONTRÔLES: {nb_ok} OK | {nb_alerte} Alertes")

        # Détail par SCPI
        print(f"\n{'─'*80}")
        print("DÉTAIL PAR SCPI")
        print(f"{'─'*80}")

        for ind in self.results['indicateurs']:
            print(f"\n🏢 {ind.nom_scpi}")
            print(f"   Société: {ind.societe_gestion or 'N/A'}")
            print(f"   TD: {ind.taux_distribution or 'N/A'}% | TOF: {ind.tof_annuel or 'N/A'}%")
            print(f"   Indicateurs: {ind.nb_indicateurs_extraits} ({ind.qualite_extraction})")


# =============================================================================
# MAIN
# =============================================================================

def main():
    import argparse

    parser = argparse.ArgumentParser(description="SCPI Consolidator V6 - Extraction Académique Avancée")
    parser.add_argument("-i", "--input", help="Dossier d'entrée", default="input")
    parser.add_argument("-o", "--output", help="Dossier de sortie", default="output")
    args = parser.parse_args()

    consolidator = SCPIConsolidatorV6(output_dir=args.output, input_dir=args.input)

    if consolidator.pdf_files:
        print(f"📄 {len(consolidator.pdf_files)} fichier(s) PDF trouvé(s)")
    else:
        print(f"⚠️  Aucun fichier PDF trouvé dans {consolidator.input_dir}")

    consolidator.consolidate()

    excel_path = consolidator.export_excel()
    csv_paths = consolidator.export_csv()

    consolidator.print_summary()

    print(f"\n📁 Fichiers générés:")
    print(f"   Excel: {excel_path}")
    for csv in csv_paths:
        print(f"   CSV: {csv}")


if __name__ == "__main__":
    main()
